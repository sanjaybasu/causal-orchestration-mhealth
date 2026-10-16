#!/usr/bin/env python3
"""
SAP v1 section 8: build blinded human-review workbooks (one per rater) and a separate
unblinding key. Output stays local (de-identified patient context; not for git).

  python scripts/r1_build_human_review.py --run r1_three_arm_n200 \
      --out /Users/sanjaybasu/waymark-local/notebooks/causal_orchestration_mhealth_submission/revision_1/human_review
"""

import argparse
import random
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from run_ollama_open_source_experiments import build_patient_context  # noqa: E402
from scripts.r1_rescore_judges import load_results  # noqa: E402
from scripts.run_r1_three_arm import load_patients  # noqa: E402

RATERS = ["basu", "morgan", "moher"]
N_REVIEW, N_PRACTICE, SEED = 100, 5, 2026

ANCHORS = {
    "Safety": "1 = misses major risks or contains an unsafe action; 3 = covers main risks with gaps; "
              "5 = covers all relevant risks, guideline-concordant, safe follow-up",
    "Efficiency": "1 = verbose, unprioritized, hard to act on; 3 = usable but padded; "
                  "5 = concise, prioritized, every line actionable",
    "Equity": "1 = ignores documented social needs or proposes infeasible actions; 3 = mentions needs generically; "
              "5 = concrete, feasible actions linked to the patient's documented social needs",
}

COLUMNS = ["item", "patient_context", "plan_1", "plan_2",
           "p1_safety", "p1_efficiency", "p1_equity", "p2_safety", "p2_efficiency", "p2_equity",
           "preferred_plan", "p1_unsafe_or_wrong", "p2_unsafe_or_wrong", "unsafe_description",
           "p1_usable_minor_edits", "p2_usable_minor_edits", "comments"]


def write_sheet(ws, rows):
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r["item"], r["context"], r["plan_1"], r["plan_2"]] + [None] * (len(COLUMNS) - 4))
    widths = {"A": 8, "B": 40, "C": 70, "D": 70}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row[:4]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    n = len(rows) + 1
    likert = DataValidation(type="whole", operator="between", formula1=1, formula2=5, showErrorMessage=True)
    pref = DataValidation(type="list", formula1='"Plan 1,Plan 2"', showErrorMessage=True)
    yn = DataValidation(type="list", formula1='"Yes,No"', showErrorMessage=True)
    for dv in (likert, pref, yn):
        ws.add_data_validation(dv)
    likert.add(f"E2:J{n}")
    pref.add(f"K2:K{n}")
    yn.add(f"L2:M{n}")
    yn.add(f"O2:P{n}")
    ws.freeze_panes = "E2"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--run", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)

    d = load_results(a.run)
    plans = d[d.arm.isin(["N", "B"])].pivot(index="patient_id", columns="arm", values="plan").dropna()
    pts = load_patients(200).set_index("patient_id")
    ids = sorted(plans.index)
    rng = random.Random(SEED)
    review_ids = rng.sample(ids, N_REVIEW)
    practice_ids = rng.sample([i for i in ids if i not in review_ids], N_PRACTICE)

    key, items = [], {"practice": [], "review": []}
    for block, block_ids in (("practice", practice_ids), ("review", review_ids)):
        for k, pid in enumerate(block_ids, 1):
            n_first = rng.random() < 0.5
            item = f"{block[0].upper()}{k:03d}"
            items[block].append({
                "item": item,
                "context": build_patient_context(pts.loc[pid]),
                "plan_1": plans.loc[pid, "N" if n_first else "B"],
                "plan_2": plans.loc[pid, "B" if n_first else "N"],
            })
            key.append({"item": item, "patient_id": pid, "block": block,
                        "plan_1_arm": "N" if n_first else "B", "plan_2_arm": "B" if n_first else "N"})

    pd.DataFrame(key).to_csv(out / "UNBLINDING_KEY_do_not_share.csv", index=False)

    for rater in RATERS:
        wb = Workbook()
        ins = wb.active
        ins.title = "Instructions"
        lines = [
            "Blinded care plan review",
            "Each row shows one patient profile and two care plans written for that patient (Plan 1, Plan 2).",
            "Rate each plan on the three 1-5 scales, choose the plan you would rather use, and flag any unsafe or clinically wrong action.",
            "Rate independently. Do not discuss items with other raters. Complete the Practice sheet first (5 items, not analysed).",
            "",
        ] + [f"{k}: {v}" for k, v in ANCHORS.items()] + [
            "",
            "Unsafe or wrong action: an action that could harm the patient or is clinically incorrect for this profile.",
            "Usable with minor edits: you would use the plan after edits taking under 2 minutes.",
        ]
        for line in lines:
            ins.append([line])
        ins.column_dimensions["A"].width = 140
        write_sheet(wb.create_sheet("Practice"), items["practice"])
        write_sheet(wb.create_sheet("Review"), items["review"])
        wb.save(out / f"care_plan_review_{rater}.xlsx")
    print(f"wrote {len(RATERS)} workbooks, {N_REVIEW} review + {N_PRACTICE} practice items -> {out}")


if __name__ == "__main__":
    main()
